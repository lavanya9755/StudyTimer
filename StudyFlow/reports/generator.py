"""
StudyFlow - Report Generator
Generates PDF and Excel reports for study sessions.
"""

import os
import tempfile
from datetime import date, datetime
from typing import List, Dict
from models.session import StudySession


def generate_pdf_report(sessions: List[StudySession], stats: Dict,
                        bar_chart_buf, pie_chart_buf) -> str:
    """Generate a PDF daily report. Returns the file path."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table,
            TableStyle, Image as RLImage, HRFlowable,
        )
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        import io

        today = date.today().strftime("%Y-%m-%d")
        path = os.path.join(os.path.expanduser("~"), f"StudyFlow_Report_{today}.pdf")

        doc = SimpleDocTemplate(path, pagesize=A4,
                                leftMargin=2*cm, rightMargin=2*cm,
                                topMargin=2*cm, bottomMargin=2*cm)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("title", fontSize=22, textColor=colors.HexColor("#c084fc"),
                                     alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=6)
        sub_style = ParagraphStyle("sub", fontSize=11, textColor=colors.HexColor("#818cf8"),
                                   alignment=TA_CENTER, spaceAfter=4)
        body_style = ParagraphStyle("body", fontSize=10, textColor=colors.HexColor("#1a1025"),
                                    spaceAfter=4)
        section_style = ParagraphStyle("section", fontSize=13, textColor=colors.HexColor("#c084fc"),
                                       fontName="Helvetica-Bold", spaceAfter=6, spaceBefore=12)

        story = []

        # Header
        story.append(Paragraph("🐱 StudyFlow Daily Report", title_style))
        story.append(Paragraph(f"Date: {today}", sub_style))
        story.append(HRFlowable(width="100%", color=colors.HexColor("#c084fc")))
        story.append(Spacer(1, 0.4*cm))

        # Stats
        story.append(Paragraph("📊 Daily Statistics", section_style))
        stat_data = [
            ["Total Study Time", stats.get("total_display", "0h 00m")],
            ["Sessions", str(stats.get("session_count", 0))],
            ["Longest Session", stats.get("longest_display", "0h 00m")],
            ["Average Session", stats.get("avg_display", "0h 00m")],
            ["Productivity Score", f"{stats.get('productivity_score', 0)}/100"],
        ]
        stat_table = Table(stat_data, colWidths=[8*cm, 8*cm])
        stat_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e9d5ff")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c084fc")),
            ("ROWBACKGROUNDS", (1, 0), (1, -1), [colors.white, colors.HexColor("#fdf4ff")]),
            ("ALIGN", (1, 0), (1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(stat_table)
        story.append(Spacer(1, 0.4*cm))

        # Summary
        story.append(Paragraph(
            f"{stats.get('summary_emoji','')} {stats.get('summary_title','')} — {stats.get('summary_body','')}",
            body_style
        ))
        story.append(Spacer(1, 0.4*cm))

        # Sessions table
        if sessions:
            story.append(Paragraph("📋 Session Log", section_style))
            session_data = [["#", "Start", "End", "Duration", "Notes", "Remark"]]
            for i, s in enumerate(sessions, 1):
                remark_short = s.remark[:40] + "..." if len(s.remark) > 40 else s.remark
                session_data.append([
                    str(i), s.start_time, s.end_time,
                    s.duration_display(), s.notes[:30] or "—", remark_short
                ])

            session_table = Table(session_data, colWidths=[1*cm, 2.5*cm, 2.5*cm, 2.5*cm, 3*cm, 4.5*cm])
            session_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#c084fc")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e9d5ff")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fdf4ff")]),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(session_table)
            story.append(Spacer(1, 0.4*cm))

        # Charts
        for buf, caption in [(bar_chart_buf, "Study Hours per Session"),
                              (pie_chart_buf, "Productivity Distribution")]:
            if buf:
                buf.seek(0)
                img = RLImage(buf, width=12*cm, height=5*cm)
                story.append(img)
                story.append(Paragraph(caption, sub_style))
                story.append(Spacer(1, 0.3*cm))

        doc.build(story)
        return path

    except Exception as e:
        return f"ERROR: {e}"


def generate_excel_report(sessions: List[StudySession], stats: Dict) -> str:
    """Generate an Excel report. Returns the file path."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        today = date.today().strftime("%Y-%m-%d")
        path = os.path.join(os.path.expanduser("~"), f"StudyFlow_Report_{today}.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily Report"

        purple = "C084FC"
        light_purple = "E9D5FF"
        dark = "1A1025"

        # Title
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = f"🐱 StudyFlow Daily Report — {today}"
        title_cell.font = Font(size=16, bold=True, color=purple)
        title_cell.alignment = Alignment(horizontal="center")

        # Stats
        ws.append([])
        ws.append(["📊 Daily Statistics"])
        ws["A3"].font = Font(size=12, bold=True, color=purple)

        stat_rows = [
            ("Total Study Time", stats.get("total_display", "")),
            ("Sessions", stats.get("session_count", 0)),
            ("Longest Session", stats.get("longest_display", "")),
            ("Average Session", stats.get("avg_display", "")),
            ("Productivity Score", f"{stats.get('productivity_score', 0)}/100"),
        ]
        for label, val in stat_rows:
            ws.append([label, val])

        ws.append([])
        # Sessions header
        headers = ["Session", "Start Time", "End Time", "Duration", "Notes", "Remark"]
        ws.append(headers)
        header_row = ws.max_row
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(header_row, col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=purple)
            cell.alignment = Alignment(horizontal="center")

        for i, s in enumerate(sessions, 1):
            ws.append([i, s.start_time, s.end_time, s.duration_display(), s.notes, s.remark])

        # Column widths
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

        wb.save(path)
        return path

    except Exception as e:
        return f"ERROR: {e}"
